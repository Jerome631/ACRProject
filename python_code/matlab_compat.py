import os
import sys
import numpy as np
import scipy.ndimage
import openpyxl
from openpyxl.utils import coordinate_to_tuple
import matplotlib.pyplot as plt
from matplotlib.path import Path

# NOTE: tkinter is intentionally NOT imported at module level.
# Headless environments (e.g. Streamlit Cloud) don't have a display or the
# _tkinter C extension installed, so importing tkinter at import-time would
# crash this whole module including the non-GUI helpers below (imcrop,
# imrotate, xlswrite, etc.) that other code needs regardless of GUI support.
# Each dialog function below imports tkinter lazily, only when actually called.

def get_tkinter_root():
    import tkinter as tk
    root = tk.Tk()
    root.withdraw()
    return root

def uigetfile(filter_str="*.*", title="Select File"):
    from tkinter import filedialog
    root = get_tkinter_root()
    # Parse MATLAB style filter e.g. '*.dcm'
    filetypes = []
    if filter_str:
        ext = filter_str.split(';')
        for e in ext:
            e_clean = e.strip('*')
            filetypes.append((f"Files ({e})", e))
    filetypes.append(("All files", "*.*"))
    
    filename = filedialog.askopenfilename(title=title, filetypes=filetypes)
    if not filename:
        return "", ""
    pathname, fname = os.path.split(filename)
    return fname, pathname + os.sep

def uiputfile(filter_str="*.*", title="Save As"):
    from tkinter import filedialog
    root = get_tkinter_root()
    filetypes = []
    if filter_str:
        ext = filter_str.split(';')
        for e in ext:
            filetypes.append((f"Files ({e})", e))
    filetypes.append(("All files", "*.*"))
    
    filename = filedialog.asksaveasfilename(title=title, filetypes=filetypes)
    if not filename:
        return "", ""
    pathname, fname = os.path.split(filename)
    return fname, pathname + os.sep

def questdlg(message, title="Question", opt1="Yes", opt2="No", default="Yes"):
    import tkinter as tk
    root = get_tkinter_root()
    # Custom simple dialog with 2 buttons to mimic MATLAB questdlg
    dialog = tk.Toplevel(root)
    dialog.title(title)
    dialog.geometry("350x120")
    dialog.resizable(False, False)
    dialog.grab_set()
    
    label = tk.Label(dialog, text=message, wraplength=320, justify="left", pady=10)
    label.pack()
    
    result = [default]
    
    def on_click(val):
        result[0] = val
        dialog.destroy()
        
    btn_frame = tk.Frame(dialog)
    btn_frame.pack(side="bottom", fill="x", pady=10)
    
    btn1 = tk.Button(btn_frame, text=opt1, width=10, command=lambda: on_click(opt1))
    btn1.pack(side="left", padx=20, expand=True)
    
    btn2 = tk.Button(btn_frame, text=opt2, width=10, command=lambda: on_click(opt2))
    btn2.pack(side="right", padx=20, expand=True)
    
    # Wait for the dialog to be destroyed
    dialog.wait_window()
    return result[0]

def inputdlg(prompts, title="Input Dialog", default_answers=None):
    import tkinter as tk
    root = get_tkinter_root()
    dialog = tk.Toplevel(root)
    dialog.title(title)
    dialog.grab_set()
    
    entries = []
    for i, p in enumerate(prompts):
        lbl = tk.Label(dialog, text=p, justify="left", anchor="w")
        lbl.pack(fill="x", padx=10, pady=(5, 0))
        entry = tk.Entry(dialog, width=40)
        if default_answers and i < len(default_answers):
            entry.insert(0, default_answers[i])
        entry.pack(fill="x", padx=10, pady=5)
        entries.append(entry)
        
    answers = []
    
    def on_ok():
        for entry in entries:
            answers.append(entry.get())
        dialog.destroy()
        
    btn = tk.Button(dialog, text="OK", width=10, command=on_ok)
    btn.pack(pady=10)
    
    dialog.wait_window()
    return answers

def msgbox(message, title="Message"):
    from tkinter import messagebox
    root = get_tkinter_root()
    messagebox.showinfo(title, message)
    return 1

def xlswrite(filename, val, sheetname='Sheet1', cell='A1'):
    try:
        if os.path.exists(filename):
            wb = openpyxl.load_workbook(filename)
        else:
            wb = openpyxl.Workbook()
            # remove default sheet
            if 'Sheet' in wb.sheetnames:
                wb.remove(wb['Sheet'])
        
        if sheetname not in wb.sheetnames:
            ws = wb.create_sheet(title=sheetname)
        else:
            ws = wb[sheetname]
            
        row, col = coordinate_to_tuple(cell)
        
        def clean_val(v):
            if hasattr(v, 'item'):
                return v.item()
            return v
            
        if isinstance(val, (list, np.ndarray)):
            val = np.asarray(val)
            if val.ndim == 1:
                for i, v in enumerate(val):
                    ws.cell(row=row, column=col+i, value=clean_val(v))
            elif val.ndim == 2:
                for r in range(val.shape[0]):
                    for c in range(val.shape[1]):
                        ws.cell(row=row+r, column=col+c, value=clean_val(val[r, c]))
        else:
            ws.cell(row=row, column=col, value=clean_val(val))
            
        wb.save(filename)
        wb.close()
        return 1
    except Exception as e:
        print(f"Error writing to Excel: {e}")
        return 0

# Image processing functions matching MATLAB behavior
def fspecial(filter_type, param=None):
    if filter_type == 'average':
        if param is None:
            param = [3, 3]
        if isinstance(param, (int, float)):
            r = c = int(param)
        else:
            r, c = int(param[0]), int(param[1])
        return np.ones((r, c), dtype=float) / (r * c)
        
    elif filter_type == 'disk':
        if param is None:
            radius = 5
        else:
            radius = float(param)
        rad_ceil = int(np.ceil(radius))
        y, x = np.ogrid[-rad_ceil:rad_ceil+1, -rad_ceil:rad_ceil+1]
        mask = x**2 + y**2 <= radius**2
        kernel = mask.astype(float)
        # In MATLAB, fspecial('disk') normalizes the kernel so sum(kernel) = 1
        total = np.sum(kernel)
        if total > 0:
            kernel /= total
        return kernel
    else:
        raise ValueError(f"Unsupported filter type: {filter_type}")

def imfilter(image, kernel, mode='constant'):
    # MATLAB's imfilter does correlation by default, with padding.
    # SciPy's correlate performs correlation.
    # For default, use mode='constant' (padding with 0).
    return scipy.ndimage.correlate(image.astype(float), kernel, mode=mode)

def imcrop(img, rect):
    # rect is [x, y, w, h] (0-based in Python)
    # MATLAB coordinates are 1-based, but we will convert standard cropped coordinates
    # to 0-based.
    # Note: in MATLAB [xmin, ymin, width, height] has inclusive bounds, so width and height
    # crop a region of size (height+1, width+1).
    x, y, w, h = int(rect[0]), int(rect[1]), int(rect[2]), int(rect[3])
    return img[y:y+h+1, x:x+w+1]

def imrotate(img, angle, method='bilinear'):
    # MATLAB's imrotate is counter-clockwise.
    # SciPy's rotate rotates counter-clockwise.
    # method can be 'nearest', 'bilinear', etc.
    order = 1 if method == 'bilinear' else 0
    return scipy.ndimage.rotate(img, angle, reshape=True, order=order)

def imresize(img, scale):
    # MATLAB's imresize defaults to bicubic, but let's use zoom with order=1 (bilinear)
    # or order=3 (bicubic) depending on needs. Order=1 is clean and standard.
    if isinstance(scale, (int, float)):
        return scipy.ndimage.zoom(img, scale, order=1)
    else:
        # scale is target size [new_h, new_w]
        zoom_factors = [scale[0]/img.shape[0], scale[1]/img.shape[1]]
        return scipy.ndimage.zoom(img, zoom_factors, order=1)

def roipoly(img, c, r):
    # c: column indices (x)
    # r: row indices (y)
    h, w = img.shape[:2]
    y, x = np.meshgrid(np.arange(h), np.arange(w), indexing='ij')
    points = np.vstack((x.flatten(), y.flatten())).T
    path = Path(np.vstack((c, r)).T)
    mask = path.contains_points(points).reshape((h, w))
    return mask

def imcontrast(img):
    # In MATLAB, imcontrast opens a contrast tool.
    # In Python, we can just show a message box and let the user press Enter or close a matplotlib plot window.
    # We will let matplotlib's interactive mode handle this.
    pass
